#!/usr/bin/env python3
"""Builds a manual row-height regression workbook."""

import io
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Font

HEIGHTS = [23.2, 0.5, 12.75, 14.2, 29.8, 409.0]


def main(output):
    path = Path(output)
    path.parent.mkdir(parents=True, exist_ok=True)
    base = io.BytesIO()
    book = Workbook()
    font = Font(name='Arial', size=11)
    book._fonts[0] = font
    book._named_styles[0].font = font
    rows = book.active
    rows.title = 'Rows'
    for index, height in enumerate(HEIGHTS, 1):
        rows.cell(index, 1, index)
        rows.cell(index, 2, f'=A{index}*2')
        rows.row_dimensions[index].height = height
    rows.row_dimensions[7].height = 23.2
    rows.row_dimensions[7].hidden = True
    rows['A7'] = 7
    rows['B7'] = '=A7*2'
    defaults = book.create_sheet('DefaultCustom')
    defaults.sheet_format.defaultRowHeight = 23.2
    defaults.sheet_format.customHeight = True
    defaults['A1'] = 9
    defaults['B1'] = '=A1*2'
    defaults['A2'] = 10
    defaults['B2'] = '=A2*2'
    defaults.row_dimensions[3].height = 12.75
    book.properties.creator = 'Ely compatibility regression'
    book.save(base)
    base.seek(0)
    with zipfile.ZipFile(base) as source, zipfile.ZipFile(path, 'w') as dest:
        for item in source.infolist():
            data = source.read(item)
            if item.filename == 'docProps/app.xml':
                doc = ET.fromstring(data)
                namespace = '{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}'
                doc.find(namespace + 'Application').text = 'Microsoft Excel'
                data = ET.tostring(doc, encoding='utf-8', xml_declaration=True)
            dest.writestr(item, data)


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit('usage: fractional-row-fixture.py output.xlsx')
    main(sys.argv[1])
