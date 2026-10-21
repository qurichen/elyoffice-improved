#!/usr/bin/env python3
"""Creates native-checked point-color controls without user data."""

from io import BytesIO
from pathlib import Path
import sys
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

C = '{http://schemas.openxmlformats.org/drawingml/2006/chart}'
EP = '{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}'
CASES = ('BarOmitted', 'BarFalse', 'BarTrue', 'BarStyled', 'BarFill',
         'PieOmitted', 'PieFalse', 'MultiOmitted')


def main(path):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.styles import Font

    book = Workbook()
    book.remove(book.active)
    font = Font(name='Arial', size=11)
    book._fonts[0] = font
    book._named_styles[0].font = font
    for label in CASES:
        sheet = book.create_sheet(label)
        sheet.append(['Category', 'Measure', 'Second'])
        for index in range(1, 4):
            sheet.append([f'Item {index}', index * 2, index * 3])
        for letter in 'ABCDEFGHIJKLM':
            sheet.column_dimensions[letter].width = 14
        chart = PieChart() if label.startswith('Pie') else BarChart()
        if isinstance(chart, BarChart):
            chart.type = 'col'
        chart.title = label
        chart.add_data(Reference(sheet, min_col=2, max_col=3 if label == 'MultiOmitted' else 2,
                                 min_row=1, max_row=4), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=4))
        chart.varyColors = False if label in ('BarFalse', 'BarFill', 'PieFalse') else None
        if label == 'BarTrue':
            chart.varyColors = True
        if label == 'BarStyled':
            chart.style = 2
        if label == 'BarFill':
            chart.series[0].graphicalProperties.solidFill = '7030A0'
        sheet.add_chart(chart, 'D2')
    buffer = BytesIO()
    book.save(buffer)
    with ZipFile(buffer) as original, ZipFile(path, 'w', ZIP_DEFLATED) as output:
        for item in original.infolist():
            data = original.read(item)
            if item.filename.startswith('xl/charts/chart') and item.filename.endswith('.xml'):
                root = ET.fromstring(data)
                chart = root.find(C + 'chart')
                title = chart.find(C + 'title')
                ET.SubElement(title, C + 'overlay', {'val': '0'})
                ET.SubElement(chart.find(C + 'legend'), C + 'overlay', {'val': '0'})
                data = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            output.writestr(item, data)


def legacy_copy(source, destination):
    assert not destination.exists(), 'refusing to replace a file'
    with ZipFile(source) as original, ZipFile(destination, 'x', ZIP_DEFLATED) as output:
        for item in original.infolist():
            data = original.read(item)
            if item.filename == 'docProps/app.xml':
                root = ET.fromstring(data)
                root.find(EP + 'Application').text = 'Microsoft Excel'
                root.find(EP + 'AppVersion').text = '12.0000'
                data = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            output.writestr(item, data)


if __name__ == '__main__':
    if len(sys.argv) == 4 and sys.argv[1] == '--legacy-copy':
        legacy_copy(Path(sys.argv[2]), Path(sys.argv[3]))
    elif len(sys.argv) == 2:
        main(Path(sys.argv[1]))
    else:
        sys.exit('usage: chart-appearance-fixture.py output.xlsx | --legacy-copy source.xlsx output.xlsx')
