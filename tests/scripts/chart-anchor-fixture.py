#!/usr/bin/env python3
"""Creates six chart-placement encodings without user data."""

from io import BytesIO
from pathlib import Path
import sys
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor, AnchorMarker, OneCellAnchor, TwoCellAnchor
from openpyxl.drawing.xdr import XDRPoint2D, XDRPositiveSize2D
from openpyxl.styles import Font


def main(path):
    book = Workbook()
    book.remove(book.active)
    font = Font(name='Arial', size=11)
    book._fonts[0] = font
    book._named_styles[0].font = font
    for label in ('CellFixed', 'CellResize', 'DefaultResize', 'PageFixed', 'TwoCellPage', 'TwoCellFixed'):
        sheet = book.create_sheet(label)
        sheet['A1'] = 'Synthetic chart placement; values 1, 2, 3'
        for index in range(2, 5):
            sheet.cell(index, 1, index - 1)
        sheet['B2'] = '=SUM(A2:A4)'
        for row in range(1, 25):
            sheet.row_dimensions[row].height = 20
        for letter in 'ABCDEFGHIJKL':
            sheet.column_dimensions[letter].width = 12
        chart = BarChart()
        chart.title = label
        chart.add_data(Reference(sheet, min_col=1, min_row=2, max_row=4))
        start = AnchorMarker(col=2, colOff=12700, row=2, rowOff=12700)
        end = AnchorMarker(col=8, colOff=25400, row=12, rowOff=25400)
        size = XDRPositiveSize2D(cx=3600000, cy=2400000)
        if label == 'CellFixed':
            chart.anchor = OneCellAnchor(_from=start, ext=size)
        elif label == 'PageFixed':
            chart.anchor = AbsoluteAnchor(pos=XDRPoint2D(x=1800000, y=600000), ext=size)
        else:
            mode = {'CellResize': 'twoCell', 'DefaultResize': None,
                    'TwoCellPage': 'absolute', 'TwoCellFixed': 'oneCell'}[label]
            chart.anchor = TwoCellAnchor(editAs=mode, _from=start, to=end)
        sheet.add_chart(chart)
    result = BytesIO()
    book.save(result)
    Path(path).write_bytes(result.getvalue())


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit('usage: chart-anchor-fixture.py output.xlsx')
    main(sys.argv[1])
